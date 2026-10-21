"""Render the lean BRD (Epic -> Feature -> Use case -> Stories) from lean/E*.json.

Usage: python3 build_lean.py TEMPLATE.docx OUTDIR
Outputs: Manih-BRD-Lean.docx, lean-clickup.csv (Epic/Feature/Story rows with Parent ID), lean-backlog.xlsx
"""
import csv
import datetime
import json
import sys
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "brd"))
import build_brd as B  # noqa: E402
from docx.shared import RGBColor  # noqa: E402

EPICS = ["E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8", "E9"]
RELEASES = [
    ("R1", "Foundation"), ("R2", "Charity registry"), ("R3", "Plan & opportunities"),
    ("R4", "Applications & decisions"), ("R5", "Contracts & projects"), ("R6", "Payments"),
    ("R7", "Insight & efficiency"), ("R8", "Advanced grants & AI"),
]
PRIO_FILL = {"Must": "E2EFDA", "Should": "FFF2CC", "Could": "EDEDED"}
CLICKUP_PRIO = {"Must": "High", "Should": "Normal", "Could": "Low"}

ROLES = [
    ("Admin", "Sets up the system: staff, roles, departments, configuration."),
    ("Grants Manager", "Owns the grant cycle: plan, opportunities, charities, reviews, projects; second payment approver."),
    ("Reviewer", "Reviews and scores applications at the assigned stage."),
    ("Committee Secretary", "Prepares committee meetings and records decisions by minute."),
    ("Project Manager", "Follows funded projects, verifies evidence; first payment approver."),
    ("Finance Officer", "Issues transfers, records bank notices, follows commitments."),
    ("CEO", "Final approver of applications and payments; oversight dashboard."),
    ("Charity delegate / staff", "Registers the charity, applies, reports progress, uploads receipts (Charity portal)."),
]


def load(src):
    epics = [json.loads((src / f"{e}.json").read_text(encoding="utf-8")) for e in EPICS if (src / f"{e}.json").exists()]
    feats = {f["id"]: (e, f) for e in epics for f in e["features"]}
    return epics, feats


def label_table(doc, rows, w=(2000, 8255)):
    t = B.new_table(doc, list(w))
    for label, value, numbered in rows:
        if not value:
            continue
        r = t.add_row()
        B.no_split(r)
        for i, c in enumerate(r.cells):
            B.set_cell_width(c, w[i])
        B.cell_text(r.cells[0], label, bold=True, color="FFFFFF", size=9.5)
        B.set_cell_fill(r.cells[0], B.TEAL)
        B.set_cell_valign(r.cells[0])
        B.cell_text(r.cells[1], value, bullets=True, numbered=numbered)
        B.set_cell_valign(r.cells[1])
    return t


def h(doc, text, level, new_page=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if new_page:
        p.paragraph_format.page_break_before = True
    return p


def word(epics, feats, template, out):
    doc = B.prepare(template)
    for part in (doc.sections[0].header, doc.sections[0].even_page_header,
                 doc.sections[0].footer, doc.sections[0].even_page_footer):
        for p in part.paragraphs:
            for r in p.runs:
                if "Business Requirements Document" in r.text:
                    r.text = r.text.replace("Business Requirements Document", "Business Requirements (Lean)")
    today = datetime.date.today().strftime("%d/%m/%Y")
    n_stories = sum(len(f["stories"]) for e in epics for f in e["features"])

    # title block
    p = doc.add_paragraph()
    B.para_fmt(p, after=40)
    B.add_run(p, "Manih  |  مانح", bold=True, color=B.TEAL, size=26)
    p = doc.add_paragraph()
    B.para_fmt(p, after=200)
    B.bottom_rule(p)
    B.add_run(p, "Grant Management System — Business Requirements (Lean)", color=B.TEAL_DARK, size=14)
    t = B.new_table(doc, [2600, 7655])
    for k, v in (("Version", "V2.0 — lean, release-based"), ("Date", today),
                 ("Structure", f"{len(epics)} epics · {len(feats)} features · {n_stories} user stories"),
                 ("Replaces", "Manih BRD V1 (96 use cases) — kept as detailed reference")):
        r = B.add_row(t, [k, v], fills={0: B.TEAL})
        r.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        r.cells[0].paragraphs[0].runs[0].bold = True

    h(doc, "1.  Document Control", 2)
    t = B.new_table(doc, [2600, 4655, 1500, 1500], header=["By", "Additions/Modifications", "Date", "Version"])
    B.add_row(t, ["Business Analysis", "Lean BRD: epics, features, use cases and user stories ordered by release.", today, "V2.0"])
    t = B.new_table(doc, [2600, 4655, 3000], header=["", "Name", "Date"])
    for k in ("Approved By", "Reviewed By"):
        r = B.add_row(t, [k, "", ""], fills={0: B.TEAL})
        r.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")

    h(doc, "2.  Purpose and how to read", 2)
    B.body(doc, "This document defines what Manih must do, organised so the team can build, test and release it step by step. "
                "Today the Foundation runs the whole grant cycle manually; each release replaces part of that manual work.")
    B.bullet_list(doc, [
        "Epic — a business capability (e.g. Payments).",
        "Feature — a releasable slice of an epic. It depends only on features released before it.",
        "Use case — one short flow per feature: who does what, and the rules.",
        "User story — a small, testable piece of a feature with Given/When/Then acceptance criteria.",
        "Ready = acceptance criteria agreed, dependencies released, no open decision. Done = all acceptance criteria pass in test.",
    ])

    h(doc, "3.  Users", 2)
    t = B.new_table(doc, [2600, 7655], header=["Role", "Responsibility"])
    for r in ROLES:
        B.add_row(t, list(r))

    h(doc, "4.  Release plan", 2)
    B.body(doc, "Build in this order. R1–R6 deliver the minimum complete grant cycle; R7–R8 add control, insight and advanced models.")
    t = B.new_table(doc, [700, 2400, 7155], header=["#", "Release", "Features (in build order)"])
    for rid, rname in RELEASES:
        fs = [f for _, f in sorted(feats.values(), key=lambda x: x[1]["id"]) if f["release"] == rid]
        B.add_row(t, [rid, rname, [f"{f['id']} {f['name']}" for f in fs]])

    h(doc, "5.  Epics and features", 2)
    t = B.new_table(doc, [700, 2300, 700, 3255, 700, 900, 600, 1100],
                    header=["Epic", "Name", "ID", "Feature", "Rel.", "Priority", "Size", "Stories"])
    for e in epics:
        for f in e["features"]:
            r = B.add_row(t, [e["epic"], e["name"], f["id"], f["name"], f["release"], f["priority"], f["size"], str(len(f["stories"]))])
            B.set_cell_fill(r.cells[5], PRIO_FILL.get(f["priority"], "EDEDED"))

    for e in epics:
        h(doc, f"{e['epic']}  {e['name']}", 1, new_page=True)
        B.body(doc, e["goal"])
        for f in e["features"]:
            h(doc, f"{f['id']}  {f['name']}", 2)
            B.body(doc, f["summary"])
            label_table(doc, [
                ("Release / priority", f"{f['release']} · {f['priority']} · size {f['size']}", False),
                ("Depends on", ", ".join(f["depends_on"]) or "None", False),
                ("Ready", f"Needs decision: {f['decision']}" if f.get("decision") else "Yes", False),
                ("Detailed BRD ref.", ", ".join(f.get("old_ucs", [])), False),
            ])
            uc = f["use_case"]
            p = doc.add_paragraph()
            B.para_fmt(p, after=60, keep_next=True)
            p.paragraph_format.space_before = B.Pt(8)
            B.add_run(p, "Use case", bold=True, color=B.TEAL_DARK, size=11)
            label_table(doc, [
                ("Actors", uc.get("actors", []), False),
                ("Trigger", uc.get("trigger", ""), False),
                ("Pre-conditions", uc.get("preconditions", []), False),
                ("Main flow", uc.get("main_flow", []), True),
                ("Alternative flows", uc.get("alternative_flows", []), False),
                ("Business rules", uc.get("business_rules", []), False),
            ])
            p = doc.add_paragraph()
            B.para_fmt(p, after=60, keep_next=True)
            p.paragraph_format.space_before = B.Pt(8)
            B.add_run(p, "User stories", bold=True, color=B.TEAL_DARK, size=11)
            t = B.new_table(doc, [1200, 3500, 5555], header=["ID", "Story", "Acceptance criteria"])
            for s in f["stories"]:
                story = [f"{s['title']}", s["story"]] + ([f"Note: {s['notes']}"] if s.get("notes") else [])
                r = B.add_row(t, [s["id"], story, s["acceptance"]])
                r.cells[1].paragraphs[0].runs[0].bold = True
            B.spacer(doc, 6)

    h(doc, "Appendix A.  Open decisions", 1, new_page=True)
    t = B.new_table(doc, [900, 3300, 6055], header=["Feature", "Name", "Decision needed before the sprint"])
    for e in epics:
        for f in e["features"]:
            if f.get("decision"):
                B.add_row(t, [f["id"], f["name"], f["decision"]])
    h(doc, "Appendix B.  Traceability to the detailed BRD", 1)
    t = B.new_table(doc, [1500, 1000, 7755], header=["Detailed BRD use case", "Feature", "Feature name"])
    rows = sorted(((u, f["id"], f["name"]) for e in epics for f in e["features"] for u in f.get("old_ucs", [])))
    for r in rows:
        B.add_row(t, list(r))
    doc.save(out)
    B.normalize(out)


def tabular(epics, outdir):
    hdr = ["Task Name", "Task Type", "Custom ID", "Parent ID", "Epic", "Release", "Priority", "Size",
           "Depends On", "Description", "Acceptance Criteria", "Ready"]
    rows = []
    for e in epics:
        rows.append([f"{e['epic']} – {e['name']}", "Epic", e["epic"], "", e["epic"], "", "", "", "", e["goal"], "", ""])
        for f in e["features"]:
            uc = f["use_case"]
            desc = "\n".join([f["summary"], "", "Main flow:"] + [f"{i + 1}. {s}" for i, s in enumerate(uc.get("main_flow", []))]
                             + ["", "Business rules:"] + [f"- {r}" for r in uc.get("business_rules", [])]
                             + ["", f"Detailed BRD: {', '.join(f.get('old_ucs', []))}"])
            rows.append([f"{f['id']} – {f['name']}", "Feature", f["id"], e["epic"], e["epic"], f["release"],
                         CLICKUP_PRIO.get(f["priority"], "Low"), f["size"], ", ".join(f["depends_on"]), desc, "",
                         "Needs decision: " + f["decision"] if f.get("decision") else "Ready"])
            for s in f["stories"]:
                rows.append([f"{s['id']} – {s['title']}", "Story", s["id"], f["id"], e["epic"], f["release"],
                             CLICKUP_PRIO.get(f["priority"], "Low"), "", "",
                             s["story"] + (f"\n\nNote: {s['notes']}" if s.get("notes") else ""),
                             "\n".join(f"{i + 1}. {a}" for i, a in enumerate(s["acceptance"])), ""])
    with open(outdir / "lean-clickup.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(hdr)
        w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Backlog"
    ws.append(hdr)
    fills = {"Epic": "4C8479", "Feature": "DCE9E5"}
    thin = Side(style="thin", color="BFBFBF")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows():
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4C8479")
    for i in range(2, ws.max_row + 1):
        typ = ws.cell(i, 2).value
        if typ in fills:
            for c in ws[i]:
                c.fill = PatternFill("solid", fgColor=fills[typ])
                c.font = Font(bold=True, color="FFFFFF" if typ == "Epic" else "000000")
    for col, wd in zip("ABCDEFGHIJKL", (48, 9, 10, 9, 6, 8, 8, 6, 12, 70, 70, 20)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(outdir / "lean-backlog.xlsx")
    return len(rows)


if __name__ == "__main__":
    template, outdir = sys.argv[1], Path(sys.argv[2])
    epics, feats = load(HERE)
    word(epics, feats, template, outdir / "Manih-BRD-Lean.docx")
    n = tabular(epics, outdir)
    print(len(epics), "epics,", len(feats), "features,", sum(len(f["stories"]) for e in epics for f in e["features"]), "stories,", n, "ClickUp rows")
