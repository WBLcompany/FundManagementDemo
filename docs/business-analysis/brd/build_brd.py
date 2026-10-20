"""Render the business-oriented BRD (Word) from brd/<CODE>.json using the client's BRD template.

Usage: python3 build_brd.py TEMPLATE.docx OUT.docx [--mmdc PATH --puppeteer CONFIG]
Diagrams are rendered with mermaid-cli when --mmdc is given; otherwise the diagram source is shown as text.
"""
import argparse
import copy
import datetime
import json
import re
import subprocess
import tempfile
from pathlib import Path

import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
ORDER = ["PLT", "CFG", "PLN", "PRG", "PTL", "BEN", "REV", "PRJ", "FIN", "TSK"]
TEAL, TEAL_DARK, GREY, ERR_FILL = "4C8479", "487E74", "F2F2F2", "FBE4D5"
TABLE_W = 10255  # dxa, as in the template
ARABIC = re.compile(r"[؀-ۿ]")
MSG_PREFIX = {"Error": "ERR", "Warning": "WRN", "Confirmation": "CNF", "Notification": "NTF", "Information": "INF"}

GLOSSARY = [
    ("BRD", "Business Requirements Document"),
    ("UC", "Use Case — an end-to-end business goal of an actor"),
    ("BRQ", "Business Requirement — a testable statement of what the system must do"),
    ("Grantor / Foundation", "The organization that funds grants and operates Manih (مؤسسة نورة الملاحي)"),
    ("Partner Organization", "A non-profit or entity that applies for and executes grants (الجهة الشريكة)"),
    ("Grant Opportunity", "A published funding call partners can apply to (فرصة منح)"),
    ("Initiative", "A plan item / program line that grants are funded from (مبادرة)"),
    ("Endowment", "A funding source (waqf) whose budget is split across plan items (وقف)"),
    ("Workflow Template / Engine", "Reusable definition of review stages, criteria and SLAs (محرك سير الإجراء)"),
    ("BOQ", "Bill of Quantities — itemized deliverables and costs of an application or contract"),
    ("Disbursement Order", "Authorization to transfer a payment to the payee (إذن صرف)"),
    ("Evidence Document", "Supporting file required for an auditable financial or plan change (مستند شاهد)"),
    ("SLA", "Service Level Agreement — the allowed days for a stage before it is overdue"),
    ("RTL", "Right-to-left layout used for Arabic"),
]

ROLES = [
    ("System Administrator", "Configures the platform, organizations, users and roles.", "Settings, user & role management, all modules"),
    ("Grants Manager", "Owns the grant cycle end to end: plans, opportunities, partners, reviews and projects.", "Plans, opportunities, applications, partners, projects, settings"),
    ("Reviewer", "Evaluates applications at the assigned stage and recommends decisions.", "Applications, partners, tasks, dashboard"),
    ("Committee Secretary", "Prepares committee meetings and records minutes and bulk decisions.", "Tasks, applications, committee minutes"),
    ("Project Manager", "Follows execution of funded projects, approves deliverables and first payment step.", "Projects, tasks"),
    ("Finance Officer", "Issues disbursement orders, records transfers and follows commitments.", "Plans, projects, payments, dashboard"),
    ("CEO / Executive Approver", "Final approver for applications and payments above the operational level.", "Dashboard, plans, applications, projects, tasks"),
    ("Partner Organization", "Registers, applies for grants, executes projects, submits evidence and receipts.", "Partner portal"),
]

DISCOVERY = [
    ("Partners & registration", "How is the partner delegate verified today? Who approves the bank account?"),
    ("Plan & opportunities", "How is the annual plan built and amended today? Which eligibility conditions are most common?"),
    ("Applications & decisions", "How many stages does an application pass through? Where are decisions usually delayed?"),
    ("Projects & evidence", "How is partner progress followed today? Who verifies evidence?"),
    ("Finance", "How long does a disbursement take today? Is there an integration with the accounting system?"),
    ("Governance & reporting", "Which reports does management request? How are tasks assigned today?"),
]


# ---------------------------------------------------------------- low-level helpers
def set_cell_fill(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_valign(cell, val="center"):
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), val)
    cell._tc.get_or_add_tcPr().append(v)


def set_cell_width(cell, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.insert(0, tcW)
    tcW.set(qn("w:w"), str(w))
    tcW.set(qn("w:type"), "dxa")


def para_fmt(p, indent=0, after=60, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.left_indent = Pt(indent)
    pf.first_line_indent = Pt(0)
    pf.space_after = Pt(after / 20 * 1) if after else Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0
    if keep_next:
        pf.keep_with_next = True
    if align is not None:
        p.alignment = align


def set_rtl(p, run):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.insert(0, bidi)
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)
    run.font.name = "Arial"
    run._r.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:cs"), "Arial")


def add_run(p, text, bold=False, color=None, size=None, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
        r._r.get_or_add_rPr()
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        r._r.rPr.append(szcs)
    if ARABIC.search(text) and not re.search(r"[A-Za-z]{3}", text):
        set_rtl(p, r)
    return r


def cell_text(cell, text, bold=False, color=None, size=9.5, bullets=False, numbered=False):
    """Fill a cell with text; a list becomes one paragraph per item (bulleted or numbered)."""
    items = text if isinstance(text, list) else [text]
    items = [i for i in items if str(i).strip()] or [""]
    first = cell.paragraphs[0]
    for k, item in enumerate(items):
        p = first if k == 0 else cell.add_paragraph()
        para_fmt(p, after=40)
        prefix = f"{k + 1}. " if numbered and len(items) > 1 else ("• " if bullets and len(items) > 1 else "")
        if prefix:
            p.paragraph_format.left_indent = Pt(12)
            p.paragraph_format.first_line_indent = Pt(-12)
        add_run(p, prefix + str(item), bold=bold, color=color, size=size)


def new_table(doc, widths, header=None, header_fill=TEAL):
    t = doc.add_table(rows=0, cols=len(widths))
    tbl = t._tbl
    tblPr = tbl.tblPr
    for tag in ("w:tblStyle",):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{e}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    t._widths = widths
    if header:
        row = add_row(t, header, fill=header_fill, bold=True, color="FFFFFF")
        trPr = row._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        trPr.append(h)
    return t


def add_row(t, values, fill=None, bold=False, color=None, fills=None):
    row = t.add_row()
    no_split(row)
    for i, (cell, v) in enumerate(zip(row.cells, values)):
        set_cell_width(cell, t._widths[i])
        cell_text(cell, v, bold=bold, color=color, bullets=True)
        f = (fills or {}).get(i, fill)
        if f:
            set_cell_fill(cell, f)
        set_cell_valign(cell)
    return row


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def merge_row(row, start, end):
    a = row.cells[start]
    a.merge(row.cells[end])
    return a


def heading(doc, text, level):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.page_break_before = True
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    para_fmt(p, indent=0 if not indent else 0, after=160)
    add_run(p, text)
    return p


def bullet_list(doc, items, numbered=False):
    for k, it in enumerate(items):
        p = doc.add_paragraph()
        para_fmt(p, after=80)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-14)
        add_run(p, (f"{k + 1}.  " if numbered else "•  ") + it)


def spacer(doc, pts=6):
    p = doc.add_paragraph()
    para_fmt(p, after=pts * 20)


def caption(doc, text):
    p = doc.add_paragraph()
    para_fmt(p, after=200)
    add_run(p, text, italic=True, size=9, color="595959")


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                f.set(qn("w:dirty"), "true")
            r._r.append(f)
        elif text:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = "Right-click and choose Update Field to build the table of contents."
            r._r.append(t)


# ---------------------------------------------------------------- diagrams
class Diagrams:
    def __init__(self, mmdc, pp):
        self.mmdc, self.pp = mmdc, pp
        self.dir = Path(tempfile.mkdtemp(prefix="brd-mmd-"))
        self.n = 0
        cfg = {"theme": "base", "themeVariables": {
            "primaryColor": "#E6F1EE", "primaryBorderColor": "#4C8479", "primaryTextColor": "#1F2D2A",
            "lineColor": "#4C8479", "secondaryColor": "#F2F2F2", "tertiaryColor": "#FFFFFF",
            "fontFamily": "Arial", "fontSize": "15px"}}
        self.cfg = self.dir / "cfg.json"
        self.cfg.write_text(json.dumps(cfg))

    def render(self, src):
        if not self.mmdc:
            return None
        self.n += 1
        i, o = self.dir / f"d{self.n}.mmd", self.dir / f"d{self.n}.png"
        i.write_text(src, encoding="utf-8")
        r = subprocess.run([self.mmdc, "-p", self.pp, "-c", str(self.cfg), "-i", str(i), "-o", str(o),
                            "-s", "2", "-b", "white", "-w", "1600"], capture_output=True, text=True)
        return o if r.returncode == 0 and o.exists() else None


def add_diagram(doc, diagrams, src, title):
    from PIL import Image
    png = diagrams.render(src)
    if png and re.match(r"\s*(flowchart|graph) LR", src):
        w, h = Image.open(png).size
        if w / h > 3.2:
            alt = diagrams.render(re.sub(r"(flowchart|graph) LR", r"\1 TD", src, count=1))
            if alt:
                png = alt
    if png:
        w, h = Image.open(png).size
        width = min(6.9, w / 2 / 96)
        height = h / 2 / 96 * (width / (w / 2 / 96))
        if height > 8.0:
            width *= 8.0 / height
        p = doc.add_paragraph()
        para_fmt(p, after=60, align=1)
        p.add_run().add_picture(str(png), width=Inches(width))
    else:
        p = doc.add_paragraph()
        para_fmt(p, after=60)
        add_run(p, src, size=8)
    caption(doc, f"Figure: {title}")


# ---------------------------------------------------------------- template preparation
def prepare(template):
    doc = docx.Document(template)
    bodyel = doc.element.body
    for el in list(bodyel):
        if el.tag != qn("w:sectPr"):
            bodyel.remove(el)
    # drop parts that belonged to the example content (comments, embedded sheet, cover picture)
    for rId, rel in list(doc.part.rels.items()):
        if any(k in rel.reltype for k in ("comments", "commentsExtended", "/package", "/image", "customXml")):
            doc.part.drop_rel(rId)
    sec = doc.sections[0]
    title = "Manih Grant Management System — Business Requirements Document"
    for hdr in (sec.header, sec.first_page_header, sec.even_page_header):
        hdr.is_linked_to_previous = False
        for el in list(hdr._element):
            hdr._element.remove(el)
    for ftr in (sec.footer, sec.first_page_footer, sec.even_page_footer):
        ftr.is_linked_to_previous = False
        for el in list(ftr._element):
            ftr._element.remove(el)
    for hdr in (sec.header, sec.even_page_header):
        p = hdr.add_paragraph()
        para_fmt(p, after=0)
        add_run(p, "Manih  |  مانح", bold=True, color=TEAL, size=12)
        p.add_run("\t")
        add_run(p, "Business Requirements Document", color="595959", size=9)
        tabs = p.paragraph_format.tab_stops
        tabs.add_tab_stop(Inches(7.1), alignment=2)
        bottom_rule(p)
    sec.first_page_header.add_paragraph()
    for ftr in (sec.footer, sec.even_page_footer):
        p = ftr.add_paragraph()
        para_fmt(p, after=0)
        top_rule(p)
        add_run(p, title, color="595959", size=8.5)
        p.add_run("\t")
        page_field(p)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), alignment=2)
    sec.first_page_footer.add_paragraph()
    # remove the template's numbering/indent from headings: we write numbers explicitly
    for lvl in (1, 2, 3):
        st = doc.styles[f"Heading {lvl}"]
        st.paragraph_format.left_indent = Pt(0)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[lvl])
        st.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[lvl])
        st.paragraph_format.line_spacing = 1.0
        pPr = st.element.get_or_add_pPr()
        num = pPr.find(qn("w:numPr"))
        if num is not None:
            pPr.remove(num)
    h2 = doc.styles["Heading 2"]
    h2.font.color.rgb = RGBColor.from_string("469E83")
    h2.font.size = Pt(15)
    h2.font.bold = True
    doc.styles["normal"].paragraph_format.left_indent = Pt(0)
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)
    return doc


def bottom_rule(p, color=TEAL):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def top_rule(p, color="BFBFBF"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:top")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    pPr.append(bdr)


def page_field(p):
    r = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        elif txt == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " PAGE "
            r._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = txt
            r._r.append(t)


# ---------------------------------------------------------------- document sections
def cover(doc, today):
    for _ in range(7):
        spacer(doc, 12)
    p = doc.add_paragraph()
    para_fmt(p, after=200, align=1)
    add_run(p, "Manih  |  مانح", bold=True, color=TEAL, size=34)
    p = doc.add_paragraph()
    para_fmt(p, after=120, align=1)
    add_run(p, "Grant Management System", bold=True, color="404040", size=20)
    p = doc.add_paragraph()
    para_fmt(p, after=600, align=1)
    bottom_rule(p)
    add_run(p, "Business Requirements Document (BRD)", color=TEAL_DARK, size=16)
    t = new_table(doc, [3000, 4000])
    t.alignment = 1
    for k, v in (("Client", "مؤسسة نورة الملاحي — Noura Al-Mallahi Foundation"), ("Prepared by", "WBL — Business Analysis"),
                 ("Version", "V1.0 (Draft for review)"), ("Date", today), ("Source", "Manih interactive prototype")):
        add_row(t, [k, v], fills={0: TEAL}, )
        t.rows[-1].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        t.rows[-1].cells[0].paragraphs[0].runs[0].bold = True


def uc_table(doc, uc, ids):
    w = [1772, 1081, 7402]
    t = new_table(doc, w)

    def label_row(label, value, numbered=False, bullets=True, fill_value=None):
        row = t.add_row()
        no_split(row)
        for i, c in enumerate(row.cells):
            set_cell_width(c, w[i])
        lab = row.cells[0]
        cell_text(lab, label, bold=True, color="FFFFFF", size=10)
        set_cell_fill(lab, TEAL)
        set_cell_valign(lab)
        val = merge_row(row, 1, 2)
        cell_text(val, value, numbered=numbered, bullets=bullets)
        if fill_value:
            set_cell_fill(val, fill_value)
        set_cell_valign(val)
        return row

    def id_rows(label, items, fill):
        first = None
        for k, (code, text) in enumerate(items):
            row = t.add_row()
            for i, c in enumerate(row.cells):
                set_cell_width(c, w[i])
            cell_text(row.cells[1], code, bold=True, size=9)
            set_cell_fill(row.cells[1], fill)
            cell_text(row.cells[2], text)
            for c in row.cells:
                set_cell_valign(c)
            if k == 0:
                first = row.cells[0]
                cell_text(first, label, bold=True, color="FFFFFF", size=10)
                set_cell_fill(first, TEAL)
                set_cell_valign(first)
            else:
                set_cell_fill(row.cells[0], TEAL)
        if first is not None and len(items) > 1:
            first.merge(t.rows[-1].cells[0])

    label_row("Title", uc["title"], bullets=False)
    label_row("Use Case Description", uc["description"], bullets=False)
    label_row("Actors", uc["actors"])
    label_row("Pre-Conditions", uc["preconditions"])
    label_row("Trigger", uc["trigger"], bullets=False)
    label_row("Scenario", uc["scenario"], numbered=True)
    if uc.get("alternative_flows"):
        label_row("Alternative Flow", uc["alternative_flows"])
    id_rows("Business Requirements", ids["brq"], GREY)
    if ids["err"]:
        id_rows("Errors", ids["err"], ERR_FILL)
    label_row("Post-Condition", uc["postconditions"], bullets=False)
    return t


def fields_table(doc, fields):
    t = new_table(doc, [2000, 2000, 1250, 1000, 1100, 2905],
                  header=["Field Name [English]", "Field Name [Arabic]", "Type", "Editable", "Mandatory", "Restrictions/Notes"])
    for f in fields:
        add_row(t, [f.get("en", ""), f.get("ar", ""), f.get("type", ""), f.get("editable", ""), f.get("mandatory", ""), f.get("notes", "")])


def build(args):
    mods = []
    for code in ORDER:
        p = Path(args.data) / f"{code}.json"
        if p.exists():
            mods.append(json.loads(p.read_text(encoding="utf-8")))
    today = datetime.date.today().strftime("%d/%m/%Y")
    diagrams = Diagrams(args.mmdc, args.puppeteer)
    doc = prepare(args.template)

    # numbering across the document
    uc_no, brq_no, err_no = 0, 0, 0
    msg_no = {}
    for m in mods:
        m["_uc"] = {}
        for uc in m["use_cases"]:
            uc_no += 1
            brq = []
            for r in uc["business_requirements"]:
                brq_no += 1
                brq.append((f"BRQ{brq_no:03d}", r))
            err = []
            for e in uc.get("errors", []):
                err_no += 1
                err.append((f"ERR{err_no:03d}", [e["name"], e["description"]] if e.get("name") else e["description"]))
            m["_uc"][uc["key"]] = {"no": f"UC{uc_no:02d}", "brq": brq, "err": err}

    cover(doc, today)
    heading(doc, "Table of Contents", 1)
    add_toc(doc)

    # 1 Document control
    heading(doc, "1.  Document Control", 1)
    heading(doc, "1.1  Document History", 2)
    t = new_table(doc, [2600, 4655, 1500, 1500], header=["By", "Additions/Modifications", "Date", "Version"])
    add_row(t, ["WBL — Business Analysis", "Initial BRD derived from the Manih prototype: scope, current and future state, use cases, business requirements, rules, KPIs, gaps and open questions.", today, "V1.0"])
    heading(doc, "1.2  Approval", 2)
    t = new_table(doc, [2600, 4655, 3000], header=["", "Name", "Date"])
    add_row(t, ["Approved By", "", ""], fills={0: TEAL})
    t.rows[-1].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
    add_row(t, ["Reviewed By", "", ""], fills={0: TEAL})
    t.rows[-1].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")

    # 2 Introduction
    heading(doc, "2.  Introduction", 1)
    heading(doc, "2.1  Document Purpose", 2)
    body(doc, "The purpose of this document is to define the business requirements of Manih (مانح), a grant management system for "
              "the Noura Al-Mallahi Foundation. It describes, in business terms, how the foundation plans its annual grant budget, "
              "publishes grant opportunities, registers and verifies partner organizations, evaluates applications, contracts and "
              "monitors funded projects, and disburses payments — and what the system must do to support each step.")
    body(doc, "The requirements were derived from a systematic review of the Manih interactive prototype and are intended to be "
              "validated with the foundation's business owners before implementation.")
    heading(doc, "2.2  Intended Audience", 2)
    bullet_list(doc, ["Business owners and executive management of the foundation — to validate scope, goals and policies.",
                      "Grants, finance and project teams — to confirm processes, rules and data.",
                      "Product owner, UI/UX designers, developers and QA — to build and test the solution."])
    heading(doc, "2.3  Approach and Conventions", 2)
    body(doc, "Each use case is described with its actors, pre-conditions, trigger, main scenario, alternative flows, business "
              "requirements (BRQ), errors and post-condition. Requirements are written as testable “The system must…” statements. "
              "Where the prototype does not yet behave as a production system should, the requirement states the target behaviour. "
              "Items that need a stakeholder decision are marked “(To be confirmed)”; analyst proposals are marked “(Recommended)”.")
    heading(doc, "2.4  Definitions and Abbreviations", 2)
    t = new_table(doc, [3000, 7255], header=["Term", "Definition"])
    for a, b in GLOSSARY:
        add_row(t, [a, b])

    # 3 Overview
    heading(doc, "3.  Overview", 1)
    heading(doc, "3.1  Project Scope", 2)
    body(doc, "Manih covers the full grant lifecycle of the foundation in one Arabic-first platform, from the annual plan to the "
              "last payment receipt. The scope is organized in the following business capability areas:")
    t = new_table(doc, [700, 3300, 6255], header=["#", "Capability area", "Business scope"])
    for k, m in enumerate(mods, 1):
        add_row(t, [str(k), m["title"], m["summary"]])
    spacer(doc)
    body(doc, "Out of scope for this document: technical architecture, infrastructure, data migration from legacy systems, and "
              "integrations not evidenced in the prototype (e.g. accounting system, national identity and registry services), which "
              "are listed as open questions where relevant.")
    heading(doc, "3.2  Project Goals", 2)
    bullet_list(doc, [
        "Single grant lifecycle: manage plans, opportunities, applications, contracts, projects and payments in one system instead of fragmented tools.",
        "Faster, transparent decisions: route every application and payment through defined stages, SLAs and approvers, with a complete decision history.",
        "Budget control: keep every commitment and payment tied to the annual plan and its endowments, so available budget is always known.",
        "Verified partners: register partner organizations against official licence data and verify delegates and bank accounts before funding.",
        "Accountability for results: link grants to indicators, deliverables and evidence, and release payments against approved milestones.",
        "Better partner experience: give partners a self-service portal to apply, track, respond and submit evidence.",
        "Assisted, not automated, judgement: use AI to pre-fill forms, check documents and summarize risk, while humans keep every decision.",
    ], numbered=True)
    heading(doc, "3.3  Stakeholders and User Roles", 2)
    t = new_table(doc, [2400, 4355, 3500], header=["Role", "Responsibility", "Main areas used"])
    for r in ROLES:
        add_row(t, list(r))

    # 4 Solution space
    heading(doc, "4.  Solution Space", 1)
    heading(doc, "4.1  Current and Future State", 2)
    heading(doc, "4.1.1  Current state", 3)
    body(doc, "The foundation's current grant process and tools were not part of the material reviewed, so the current state is "
              "to be confirmed during the analysis workshop with the foundation. The prototype's workshop agenda compares each step "
              "of the grant journey with the existing process; the discovery questions below are to be answered for each area.")
    t = new_table(doc, [3000, 7255], header=["Area", "Current-state discovery questions (To be confirmed)"])
    for a, q in DISCOVERY:
        add_row(t, [a, q])
    heading(doc, "4.1.2  Future state", 3)
    body(doc, "In the future state, the whole grant journey runs in Manih. The annual plan and its endowments define the budget "
              "ceilings; grant opportunities are published against plan initiatives with clear eligibility, forms and workflow "
              "templates; partner organizations register once, are verified, and apply through the portal; applications move "
              "through staged reviews and committee decisions with full history; approved applications become contracts and "
              "projects with activities, indicators and a payment schedule; and each payment passes an approval chain, a "
              "disbursement order, a bank transfer and a partner receipt.")
    body(doc, "Every financial or plan change is supported by an evidence document, every decision is attributable to a person, "
              "and management sees live commitments, payments and performance. AI assistance proposes values and flags risks, "
              "but never takes a decision on its own.")
    heading(doc, "4.1.3  End-to-end business process", 3)
    add_diagram(doc, diagrams, E2E, "End-to-end grant lifecycle in Manih")

    section_no = 2
    for m in mods:
        heading(doc, f"4.{section_no}  Use Cases — {m['title']}", 2)
        body(doc, m["summary"])
        if m.get("process", {}).get("mermaid"):
            add_diagram(doc, diagrams, m["process"]["mermaid"], m["process"].get("title", m["title"]))
        for k, uc in enumerate(m["use_cases"], 1):
            ids = m["_uc"][uc["key"]]
            heading(doc, f"4.{section_no}.{k}  {ids['no']} – {uc['name']}", 3)
            uc_table(doc, uc, ids)
            if uc.get("fields"):
                p = doc.add_paragraph()
                para_fmt(p, after=80, keep_next=True)
                p.paragraph_format.space_before = Pt(10)
                add_run(p, f"Fields description — {ids['no']}", bold=True, color=TEAL_DARK, size=11)
                fields_table(doc, uc["fields"])
            spacer(doc, 8)
        section_no += 1

    # messages
    heading(doc, f"4.{section_no}  Messages and Notifications", 2)
    heading(doc, f"4.{section_no}.1  System notifications", 3)
    body(doc, "User-facing messages observed in the prototype (Arabic text as shown to users) and recommended where missing.")
    t = new_table(doc, [1150, 1450, 1150, 3250, 3255], header=["Code", "Type", "Use case", "Message text “English”", "Message text “Arabic”"])
    for m in mods:
        for msg in m.get("messages", []):
            typ = msg.get("type", "Information")
            pre = MSG_PREFIX.get(typ, "MSG")
            msg_no[pre] = msg_no.get(pre, 0) + 1
            uc = m["_uc"].get(msg.get("use_case"), {}).get("no", "")
            add_row(t, [f"{pre}{msg_no[pre]:03d}", typ, uc, msg.get("en", ""), msg.get("ar", "")])

    # 5 Business rules
    heading(doc, "5.  Business Rules", 1)
    body(doc, "Key policy rules the system must enforce, grouped by capability area.")
    br_no = 0
    for m in mods:
        if not m.get("business_rules"):
            continue
        heading(doc, f"5.{ORDER.index(m['module']) + 1}  {m['title']}" if False else f"{m['title']}", 3)
        t = new_table(doc, [1100, 7855, 1300], header=["ID", "Business rule", "Use case"])
        for r in m["business_rules"]:
            br_no += 1
            add_row(t, [f"BR{br_no:03d}", r["rule"], m["_uc"].get(r.get("use_case"), {}).get("no", "")])

    # 6 KPIs
    kpis = [(m["title"], k) for m in mods for k in m.get("kpis", [])]
    heading(doc, "6.  Reports and KPIs", 1)
    body(doc, "Measures and report figures the system must calculate consistently across all screens and reports.")
    t = new_table(doc, [2600, 2800, 4855], header=["Area", "KPI / figure", "Business definition"])
    for area, k in kpis:
        add_row(t, [area, k["name"], k["definition"]])

    # 7 Gaps
    heading(doc, "7.  Gaps and Recommendations", 1)
    body(doc, "Areas where the prototype does not yet meet the needs of a production grant system. Each gap is reflected as a target "
              "requirement in the use cases above.")
    t = new_table(doc, [850, 1700, 2900, 2200, 1900, 705], header=["ID", "Area", "Gap", "Business impact", "Recommendation", "Priority"])
    g = 0
    rank = {"High": 0, "Medium": 1, "Low": 2}
    allgaps = sorted([(m, x) for m in mods for x in m.get("gaps", [])], key=lambda mx: rank.get(mx[1].get("priority"), 3))
    for m, x in allgaps:
        g += 1
        fill = {"High": ERR_FILL}.get(x.get("priority"))
        add_row(t, [f"GAP{g:02d}", f"{m['title']} — {x.get('area', '')}", x["gap"], x.get("impact", ""), x.get("recommendation", ""), x.get("priority", "")],
                fills={5: fill} if fill else None)

    # 8 Open questions
    heading(doc, "8.  Open Questions", 1)
    body(doc, "Decisions required from the foundation's business owners before the requirements can be finalized.")
    t = new_table(doc, [900, 2600, 6755], header=["#", "Area", "Question"])
    q = 0
    for m in mods:
        for x in m.get("open_questions", []):
            q += 1
            add_row(t, [f"Q{q:02d}", m["title"], x])

    # Appendix traceability
    heading(doc, "Appendix A.  Traceability to the Detailed Analysis", 1)
    body(doc, "Each use case maps to the user stories (US) and functional requirements (FR) of the detailed technical analysis package "
              "(Manih-BA-Package), which remains the reference for developers and QA.")
    t = new_table(doc, [1100, 3700, 5455], header=["Use case", "Name", "Detailed analysis references"])
    for m in mods:
        for uc in m["use_cases"]:
            add_row(t, [m["_uc"][uc["key"]]["no"], uc["name"], ", ".join(uc.get("trace", []))])

    doc.save(args.out)
    print(f"{uc_no} use cases, {brq_no} business requirements, {err_no} errors, {sum(msg_no.values())} messages, {br_no} rules, {g} gaps, {q} questions, {diagrams.n} diagrams")


E2E = """flowchart TD
  A["Annual plan and endowments"] --> B["Publish grant opportunity"]
  R["Partner registers and is verified"] --> C
  B --> C["Partner submits application"]
  C --> D["Staged review and evaluation"]
  D --> E["Committee and final approval"]
  D -->|"Edits requested"| C
  E -->|"Approved"| F["Contract and project created"]
  F --> G["Execution: activities, indicators, evidence"]
  G --> H["Milestone approved, payment due"]
  H --> I["Payment approvals and disbursement order"]
  I --> J["Bank transfer and partner receipt"]
  J --> K["Project closure and reporting"]
"""

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("out")
    ap.add_argument("--mmdc")
    ap.add_argument("--puppeteer")
    ap.add_argument("--data", default=str(HERE))
    build(ap.parse_args())
