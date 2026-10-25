"""Build the BRD ↔ lean features ↔ ClickUp (S002) mapping workbook.

Usage: python3 build_mapping.py DATA_DIR OUT.xlsx
DATA_DIR holds uc_list.json, lean_features.json and mapping.json.
"""
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TEAL = "4C8479"
FILL = {"Full": "E2EFDA", "Partial": "FFF2CC", "Not in ClickUp": "F8CBAD", "ClickUp only": "DDEBF7"}
STATUS_FILL = {"approved": "E2EFDA", "pending": "FFF2CC", "draft": "EDEDED"}
THIN = Side(style="thin", color="BFBFBF")


def style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
    for row in ws.iter_rows():
        for c in row:
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def fill(cell, color):
    if color:
        cell.fill = PatternFill("solid", fgColor=color)


def main(data, out):
    ucs = json.loads((data / "uc_list.json").read_text(encoding="utf-8"))
    feats = json.loads((data / "lean_features.json").read_text(encoding="utf-8"))
    mp = json.loads((data / "mapping.json").read_text(encoding="utf-8"))
    lean_of = defaultdict(list)
    for f in feats:
        for u in f["ucs"]:
            lean_of[u].append(f"{f['id']} {f['name']}")
    by_uc = {m["uc"]: m for m in mp["uc_map"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "UC → ClickUp"
    ws.append(["UC", "Use case", "BRD section", "Capability area", "UC status", "Lean feature",
               "ClickUp feature", "ClickUp epic", "Coverage", "Notes"])
    cover = {}
    for u in ucs:
        m = by_uc.get(u["uc"], {"clickup": []})
        links = m.get("clickup", [])
        if links:
            cov = "Full" if any(l["coverage"] == "Full" for l in links) else "Partial"
        else:
            cov = "Not in ClickUp"
        cover[u["uc"]] = cov
        notes = [f"{l['feature'].split(' · ')[0]}: {l['note']}" for l in links if l.get("note")]
        if m.get("gap_note"):
            notes.append(m["gap_note"])
        ws.append([u["uc"], u["name"], u["section"], u["area"], u["status"], "\n".join(lean_of[u["uc"]]),
                   "\n".join(f"{l['feature']} ({l['coverage']})" for l in links) or "—",
                   "\n".join(sorted({l.get("epic", "") for l in links})) or "—", cov, "\n".join(notes)])
        r = ws.max_row
        fill(ws.cell(r, 5), STATUS_FILL.get(u["status"]))
        fill(ws.cell(r, 9), FILL[cov])
    style(ws, (7, 40, 9, 26, 10, 30, 46, 30, 13, 60))

    ws2 = wb.create_sheet("ClickUp → UC")
    ws2.append(["ClickUp feature", "Epic", "ClickUp status", "BRD use cases", "Mapping", "Notes"])
    for c in mp["clickup_map"]:
        kind = "Mapped" if c["ucs"] else "ClickUp only"
        ws2.append([c["feature"], c.get("epic", ""), c.get("status", ""), ", ".join(c["ucs"]) or "—", kind, c.get("note", "")])
        fill(ws2.cell(ws2.max_row, 5), FILL["ClickUp only"] if not c["ucs"] else FILL["Full"])
    style(ws2, (50, 34, 12, 26, 14, 70))

    ws3 = wb.create_sheet("Summary")
    ws3.append(["Capability area", "Use cases", "Full", "Partial", "Not in ClickUp", "Covered %"])
    areas = defaultdict(Counter)
    for u in ucs:
        areas[u["area"]][cover[u["uc"]]] += 1
    tot = Counter()
    for a, c in areas.items():
        n = sum(c.values())
        tot.update(c)
        ws3.append([a, n, c["Full"], c["Partial"], c["Not in ClickUp"], f"{round(100 * (c['Full'] + c['Partial']) / n)}%"])
    n = sum(tot.values())
    ws3.append(["Total", n, tot["Full"], tot["Partial"], tot["Not in ClickUp"], f"{round(100 * (tot['Full'] + tot['Partial']) / n)}%"])
    for c in ws3[ws3.max_row]:
        c.font = Font(bold=True)
    only = sum(1 for c in mp["clickup_map"] if not c["ucs"])
    ws3.append([])
    ws3.append(["ClickUp features (S002)", len(mp["clickup_map"])])
    ws3.append(["ClickUp-only features (no BRD use case)", only])
    style(ws3, (44, 11, 9, 9, 15, 11))
    ws3.freeze_panes = None
    wb.save(out)
    print(dict(tot), "clickup-only:", only)


if __name__ == "__main__":
    main(Path(sys.argv[1]), sys.argv[2])
