"""Render the release backlog: Word document (client template), ClickUp CSV and Excel.

Usage: python3 build_backlog.py TEMPLATE.docx OUTDIR
"""
import csv
import datetime
import re
import sys
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "brd"))
import build_brd as B  # noqa: E402
from backlog_data import FEATURES, RELEASES  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

PRIO_FILL = {"Must": "E2EFDA", "Should": "FFF2CC", "Could": "EDEDED"}
CLICKUP_PRIO = {"Must": "High", "Should": "Normal", "Could": "Low"}


def load():
    """Number features F01.. in order and resolve dependency / reference keys."""
    keys = [f[0] for f in FEATURES]
    num = {k: f"F{i + 1:02d}" for i, k in enumerate(keys)}
    ref = re.compile(r"\bF\d{2}B?\b")
    out = []
    for f in FEATURES:
        k, rel, name, value, scope, later, acc, ucs, deps, prio, size, decision = f
        out.append(dict(id=num[k], release=rel, name=name, value=value, scope=scope,
                        later=ref.sub(lambda m: num.get(m.group(0), m.group(0)), later),
                        acceptance=acc, ucs=ucs, deps=[num[d] for d in deps],
                        priority=prio.split(" ")[0] if prio.startswith("Won") is False else "Won't",
                        priority_label=prio, size=size, decision=decision))
    return out


def prio_key(p):
    return p if p in PRIO_FILL else "Could"


def word(feats, template, out):
    doc = B.prepare(template)
    for part in (doc.sections[0].header, doc.sections[0].even_page_header):
        for p in part.paragraphs:
            for r in p.runs:
                if r.text == "Business Requirements Document":
                    r.text = "Release Backlog"
    for p in (doc.sections[0].footer, doc.sections[0].even_page_footer):
        for para in p.paragraphs:
            for r in para.runs:
                if "Business Requirements Document" in r.text:
                    r.text = r.text.replace("Business Requirements Document", "Release Backlog")

    p = doc.add_paragraph()
    B.para_fmt(p, after=60)
    B.add_run(p, "Manih  |  مانح", bold=True, color=B.TEAL, size=24)
    p = doc.add_paragraph()
    B.para_fmt(p, after=200)
    B.bottom_rule(p)
    B.add_run(p, "Release Backlog — features in priority order", color=B.TEAL_DARK, size=14)
    t = B.new_table(doc, [2600, 7655])
    for k, v in (("Source", "Manih Business Requirements Document (96 use cases)"),
                 ("Features", f"{len(feats)} features in {len(RELEASES)} releases"),
                 ("Date", datetime.date.today().strftime("%d/%m/%Y")),
                 ("Status", "Draft for review")):
        r = B.add_row(t, [k, v], fills={0: B.TEAL})
        r.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        r.cells[0].paragraphs[0].runs[0].bold = True

    B.spacer(doc, 6)
    h = doc.add_paragraph("1.  How this backlog works", style="Heading 2")
    B.bullet_list(doc, [
        "Each feature is a vertical slice: when it is done it can be released and used on its own.",
        "A feature only depends on features in earlier releases (or earlier in the same release). Nothing waits for a later feature; anything that would is listed under “Later” and moved to the feature that delivers it.",
        "Releases are in priority order. R1–R6 make the minimum end-to-end grant cycle (Must); R7–R8 add insight, efficiency and advanced models (Should / Could).",
        "Size: S ≈ up to one sprint for one team; M ≈ one to two sprints. Larger items were split.",
        "Detail stays in the BRD: each feature lists the use cases (UC) it delivers, for full rules, fields and messages.",
    ])
    doc.add_paragraph("2.  The 5R check", style="Heading 2")
    t = B.new_table(doc, [1700, 8555], header=["Principle", "How the backlog meets it"])
    for a, b in (("Readable", "One short card per feature: value, scope, what is not included, acceptance criteria. Plain business language."),
                 ("Realistic", "Built on what the prototype already shows working; production gaps (security, enforced roles) are inside the first features, not added later."),
                 ("Relevant", "Every feature states the value it gives a user or the business; the releases replace the manual process step by step."),
                 ("Right-sized", "Every feature is S or M; nothing larger than two sprints."),
                 ("Ready", "Each feature has testable acceptance criteria and only released dependencies. Features with an open business decision are flagged below — the decision must be taken before the sprint.")):
        B.add_row(t, [a, b])

    doc.add_paragraph("3.  Release roadmap", style="Heading 2")
    t = B.new_table(doc, [700, 2200, 4655, 2700], header=["#", "Release", "Goal", "Features"])
    for rid, rname, goal in RELEASES:
        fs = [f["id"] for f in feats if f["release"] == rid]
        B.add_row(t, [rid, rname, goal, ", ".join(fs)])

    doc.add_paragraph("4.  Feature list", style="Heading 2")
    t = B.new_table(doc, [700, 3300, 650, 900, 600, 1500, 2605],
                    header=["ID", "Feature", "Rel.", "Priority", "Size", "Depends on", "BRD use cases"])
    for f in feats:
        r = B.add_row(t, [f["id"], f["name"], f["release"], f["priority_label"], f["size"],
                          ", ".join(f["deps"]) or "—", ", ".join(f["ucs"])])
        B.set_cell_fill(r.cells[3], PRIO_FILL.get(prio_key(f["priority"]), "EDEDED"))

    decided = [f for f in feats if f["decision"]]
    doc.add_paragraph("5.  Decisions needed before the sprint", style="Heading 2")
    B.body(doc, "These features are ready except for one business decision. Taking it before sprint planning makes them fully Ready.")
    t = B.new_table(doc, [700, 3300, 6255], header=["ID", "Feature", "Decision needed"])
    for f in decided:
        B.add_row(t, [f["id"], f["name"], f["decision"]])

    for rid, rname, goal in RELEASES:
        doc.add_paragraph(f"{rid} — {rname}", style="Heading 1").paragraph_format.page_break_before = True
        B.body(doc, goal)
        for f in [x for x in feats if x["release"] == rid]:
            doc.add_paragraph(f"{f['id']}  {f['name']}", style="Heading 3")
            card(doc, f)
    doc.save(out)
    B.normalize(out)


def card(doc, f):
    w = [1772, 8483]
    t = B.new_table(doc, w)

    def row(label, value, numbered=False):
        r = t.add_row()
        B.no_split(r)
        for i, c in enumerate(r.cells):
            B.set_cell_width(c, w[i])
        B.cell_text(r.cells[0], label, bold=True, color="FFFFFF", size=9.5)
        B.set_cell_fill(r.cells[0], B.TEAL)
        B.set_cell_valign(r.cells[0])
        B.cell_text(r.cells[1], value, bullets=True, numbered=numbered)
        B.set_cell_valign(r.cells[1])

    row("Value", f["value"])
    row("Scope", f["scope"])
    if f["later"]:
        row("Not in this feature", f["later"])
    row("Acceptance criteria", f["acceptance"], numbered=True)
    row("Depends on", ", ".join(f["deps"]) or "None — can start first")
    row("Priority / size", f"{f['priority_label']}  ·  {f['size']}")
    row("BRD use cases", ", ".join(f["ucs"]))
    row("Ready", f"Needs decision: {f['decision']}" if f["decision"] else "Yes")
    B.spacer(doc, 6)


def description(f):
    parts = [f"Value: {f['value']}", "", "Scope:"] + [f"- {s}" for s in f["scope"]]
    if f["later"]:
        parts += ["", f"Not in this feature: {f['later']}"]
    parts += ["", f"Depends on: {', '.join(f['deps']) or 'None'}", f"BRD use cases: {', '.join(f['ucs'])}"]
    if f["decision"]:
        parts += [f"Decision needed: {f['decision']}"]
    return "\n".join(parts)


def tabular(feats, outdir):
    hdr = ["Task Name", "Custom ID", "Release", "Priority", "Size", "Depends On", "Tags", "Description",
           "Acceptance Criteria", "BRD Use Cases", "Ready"]
    rows = []
    for f in feats:
        rel = next(r for r in RELEASES if r[0] == f["release"])
        rows.append([f"{f['id']} – {f['name']}", f["id"], f"{rel[0]} {rel[1]}",
                     CLICKUP_PRIO.get(f["priority"], "Low"), f["size"], ", ".join(f["deps"]),
                     f"{rel[0]},{f['priority_label']},size-{f['size']}", description(f),
                     "\n".join(f"{i + 1}. {a}" for i, a in enumerate(f["acceptance"])),
                     ", ".join(f["ucs"]), "Needs decision" if f["decision"] else "Ready"])
    with open(outdir / "backlog-clickup.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(hdr)
        w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Backlog"
    ws.append(hdr)
    for r in rows:
        ws.append(r)
    thin = Side(style="thin", color="BFBFBF")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=B.TEAL)
    for row in ws.iter_rows():
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, wd in zip("ABCDEFGHIJK", (44, 9, 22, 9, 6, 14, 22, 70, 70, 22, 14)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(outdir / "backlog.xlsx")


if __name__ == "__main__":
    template, outdir = sys.argv[1], Path(sys.argv[2])
    feats = load()
    word(feats, template, outdir / "Manih-Release-Backlog.docx")
    tabular(feats, outdir)
    print(len(feats), "features")
